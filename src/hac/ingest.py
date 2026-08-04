"""Turn an arbitrary folder of documents into normalized Doc records.

Two representations are kept for every file:
  * text, split into pages or page-sized chunks, for retrieval and cheap prompts
  * the original bytes, for PDFs and images, so the model can read scans and
    complex tables natively instead of trusting an OCR layer that may not exist
"""

from __future__ import annotations

import base64
import csv
import io
import json
import re
from pathlib import Path

from .schema import Doc, Page

PDF_EXT = {".pdf"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls"}
DOCX_EXT = {".docx"}
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm"}

IMG_MEDIA = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

CHUNK = 3000


def _slug(path: Path, root: Path) -> str:
    rel = path.relative_to(root).as_posix()
    return re.sub(r"[^A-Za-z0-9._/-]+", "_", rel)


def _chunk_pages(text: str, size: int = CHUNK) -> list[Page]:
    text = text.strip()
    if not text:
        return []
    return [Page(n=i + 1, text=text[i * size:(i + 1) * size])
            for i in range((len(text) + size - 1) // size)]


def _read_pdf(path: Path) -> tuple[list[Page], int]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return [], 1
    try:
        reader = PdfReader(str(path))
        pages = []
        for i, p in enumerate(reader.pages, start=1):
            try:
                pages.append(Page(n=i, text=(p.extract_text() or "").strip()))
            except Exception:
                pages.append(Page(n=i, text=""))
        return pages, len(reader.pages)
    except Exception:
        return [], 1


def _read_sheet(path: Path) -> list[Page]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception:
        return []
    pages: list[Page] = []
    for i, ws in enumerate(wb.worksheets, start=1):
        buf = io.StringIO()
        buf.write(f"# sheet: {ws.title}\n")
        w = csv.writer(buf, delimiter="\t", lineterminator="\n")
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                w.writerow(cells)
        pages.append(Page(n=i, text=buf.getvalue()))
    wb.close()
    return pages


def _read_docx(path: Path) -> list[Page]:
    try:
        import docx  # python-docx
    except ImportError:
        return []
    try:
        d = docx.Document(str(path))
    except Exception:
        return []
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        rows = ["\t".join(c.text.strip() for c in r.cells) for r in t.rows]
        parts.append("# table\n" + "\n".join(rows))
    return _chunk_pages("\n".join(parts))


def _read_text(path: Path) -> list[Page]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() == ".json":
        try:
            raw = json.dumps(json.loads(raw), indent=1, ensure_ascii=False)
        except Exception:
            pass
    return _chunk_pages(raw)


def load_doc(path: Path, root: Path) -> Doc | None:
    ext = path.suffix.lower()
    doc_id = _slug(path, root)
    title = path.name
    size = path.stat().st_size

    if ext in PDF_EXT:
        pages, n = _read_pdf(path)
        return Doc(id=doc_id, path=str(path), kind="pdf", title=title,
                   media_type="application/pdf", n_pages=n or max(1, len(pages)),
                   pages=pages, bytes_len=size)
    if ext in IMG_EXT:
        return Doc(id=doc_id, path=str(path), kind="image", title=title,
                   media_type=IMG_MEDIA[ext], n_pages=1, pages=[], bytes_len=size)
    if ext in SHEET_EXT:
        pages = _read_sheet(path)
        return Doc(id=doc_id, path=str(path), kind="sheet", title=title,
                   n_pages=max(1, len(pages)), pages=pages, bytes_len=size)
    if ext in DOCX_EXT:
        pages = _read_docx(path)
        return Doc(id=doc_id, path=str(path), kind="docx", title=title,
                   n_pages=max(1, len(pages)), pages=pages, bytes_len=size)
    if ext in TEXT_EXT:
        pages = _read_text(path)
        return Doc(id=doc_id, path=str(path), kind="text", title=title,
                   n_pages=max(1, len(pages)), pages=pages, bytes_len=size)
    return None


def load_corpus(docs_dir: str | Path) -> list[Doc]:
    root = Path(docs_dir)
    if not root.exists():
        raise FileNotFoundError(f"documents directory not found: {root}")
    docs: list[Doc] = []
    for p in sorted(root.rglob("*")):
        if p.is_file() and not p.name.startswith("."):
            d = load_doc(p, root)
            if d is not None:
                docs.append(d)
    return docs


# --------------------------------------------------------------- content blocks

def _b64(path: str) -> str:
    return base64.standard_b64encode(Path(path).read_bytes()).decode("ascii")


def doc_blocks(doc: Doc, *, citations: bool = True, text_char_cap: int = 60_000) -> list[dict]:
    """Content blocks for one document, ready to drop into a user message."""
    if doc.kind == "pdf":
        block: dict = {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(doc.path)},
            "title": f"{doc.id}",
        }
        if citations:
            block["citations"] = {"enabled": True}
        return [block]

    if doc.kind == "image":
        return [
            {"type": "text", "text": f"<<< document {doc.id} ({doc.title}) >>>"},
            {"type": "image",
             "source": {"type": "base64", "media_type": doc.media_type, "data": _b64(doc.path)}},
        ]

    text = doc.plain_text(text_char_cap)
    block = {
        "type": "document",
        "source": {"type": "text", "media_type": "text/plain", "data": text},
        "title": f"{doc.id}",
    }
    if citations:
        block["citations"] = {"enabled": True}
    return [block]
